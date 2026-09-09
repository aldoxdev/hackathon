from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import GastoFijo, Insumo, Receta, RecetaInsumo, Venta

MAGNITUDES_VALIDAS = {"masa", "volumen", "pieza"}
TIPOS_USO_VALIDOS = {"directo", "indirecto"}
MEDIOS_PAGO_VALIDOS = {"efectivo", "banco"}
TASA_IVA_MAP = {
    "iva_16": "iva_16", "iva 16": "iva_16", "iva16": "iva_16", "16%": "iva_16", "16": "iva_16",
    "iva_0": "iva_0", "iva 0": "iva_0", "iva0": "iva_0", "0%": "iva_0", "0": "iva_0",
    "exento": "exento",
    "no_objeto": "no_objeto", "no objeto": "no_objeto", "no objeto de impuesto": "no_objeto",
}


def _norm(value) -> str:
    return str(value).strip().lower() if value is not None else ""


def _header_map(sheet) -> dict[str, int]:
    headers: dict[str, int] = {}
    for idx, cell in enumerate(sheet[1], start=1):
        if cell.value is not None:
            headers[_norm(cell.value)] = idx
    return headers


def _get(row, headers: dict[str, int], *names: str):
    for name in names:
        idx = headers.get(name)
        if idx is not None:
            return row[idx - 1].value
    return None


def _fila_vacia(row) -> bool:
    return all(cell.value is None for cell in row)


def _parse_tasa_iva(value) -> str:
    key = _norm(value)
    if key not in TASA_IVA_MAP:
        raise ValueError(f"tasa de IVA invalida: '{value}'")
    return TASA_IVA_MAP[key]


def _parse_fecha(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        return datetime.strptime(value.strip(), "%Y-%m-%d").date()
    raise ValueError(f"fecha invalida: '{value}'")


def _parse_periodo(value) -> date:
    """Acepta 'AAAA-MM', 'AAAA-MM-DD' o una fecha/datetime; siempre normaliza al dia 1 del mes."""
    if isinstance(value, datetime):
        return date(value.year, value.month, 1)
    if isinstance(value, date):
        return date(value.year, value.month, 1)
    if isinstance(value, str) and value.strip():
        partes = value.strip().split("-")
        if len(partes) >= 2:
            return date(int(partes[0]), int(partes[1]), 1)
    raise ValueError(f"mes invalido: '{value}' (usa AAAA-MM)")


def import_excel(db: Session, file_bytes: bytes) -> dict:
    wb = load_workbook(BytesIO(file_bytes), data_only=True)

    resultado = {
        "insumos_creados": 0,
        "recetas_creadas": 0,
        "receta_insumos_creados": 0,
        "ventas_creadas": 0,
        "gastos_fijos_creados": 0,
        "errores": [],
    }

    def error(hoja: str, fila: int, mensaje: str) -> None:
        resultado["errores"].append({"hoja": hoja, "fila": fila, "mensaje": mensaje})

    insumos_por_nombre = {i.nombre.strip().lower(): i for i in db.query(Insumo).all()}
    recetas_por_nombre = {r.nombre.strip().lower(): r for r in db.query(Receta).all()}

    if "Insumos" in wb.sheetnames:
        sheet = wb["Insumos"]
        headers = _header_map(sheet)
        for fila_num, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            if _fila_vacia(row):
                continue
            try:
                nombre = _get(row, headers, "nombre")
                magnitud = _norm(_get(row, headers, "magnitud"))
                tipo_uso = _norm(_get(row, headers, "tipo de uso", "tipo_uso"))
                costo = _get(row, headers, "costo por unidad", "costo_por_unidad_base")

                if not nombre:
                    raise ValueError("falta el nombre")
                if magnitud not in MAGNITUDES_VALIDAS:
                    raise ValueError(f"magnitud invalida: '{magnitud}' (usa masa, volumen o pieza)")
                if tipo_uso not in TIPOS_USO_VALIDOS:
                    raise ValueError(f"tipo de uso invalido: '{tipo_uso}' (usa directo o indirecto)")
                if costo is None:
                    raise ValueError("falta el costo por unidad")

                nombre = str(nombre).strip()
                key = nombre.lower()
                if key in insumos_por_nombre:
                    insumo = insumos_por_nombre[key]
                    insumo.magnitud = magnitud
                    insumo.tipo_uso = tipo_uso
                    insumo.costo_por_unidad_base = float(costo)
                else:
                    insumo = Insumo(
                        nombre=nombre, magnitud=magnitud, tipo_uso=tipo_uso, costo_por_unidad_base=float(costo)
                    )
                    db.add(insumo)
                    db.flush()
                    insumos_por_nombre[key] = insumo
                    resultado["insumos_creados"] += 1
            except Exception as exc:
                error("Insumos", fila_num, str(exc))

    if "Recetas" in wb.sheetnames:
        sheet = wb["Recetas"]
        headers = _header_map(sheet)
        for fila_num, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            if _fila_vacia(row):
                continue
            try:
                nombre = _get(row, headers, "nombre")
                precio = _get(row, headers, "precio de venta", "precio_venta")
                iva_raw = _get(row, headers, "iva", "tasa_iva", "tasa de iva")

                if not nombre:
                    raise ValueError("falta el nombre")
                if precio is None:
                    raise ValueError("falta el precio de venta")
                tasa_iva = _parse_tasa_iva(iva_raw)

                nombre = str(nombre).strip()
                key = nombre.lower()
                if key in recetas_por_nombre:
                    receta = recetas_por_nombre[key]
                    receta.precio_venta = float(precio)
                    receta.tasa_iva = tasa_iva
                else:
                    receta = Receta(nombre=nombre, precio_venta=float(precio), tasa_iva=tasa_iva, origen="excel")
                    db.add(receta)
                    db.flush()
                    recetas_por_nombre[key] = receta
                    resultado["recetas_creadas"] += 1
            except Exception as exc:
                error("Recetas", fila_num, str(exc))

    if "Receta_Insumos" in wb.sheetnames:
        sheet = wb["Receta_Insumos"]
        headers = _header_map(sheet)
        for fila_num, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            if _fila_vacia(row):
                continue
            try:
                receta_nombre = _get(row, headers, "receta")
                insumo_nombre = _get(row, headers, "insumo")
                cantidad = _get(row, headers, "cantidad usada", "cantidad_usada")

                if not receta_nombre or not insumo_nombre:
                    raise ValueError("falta la receta o el insumo")
                receta = recetas_por_nombre.get(_norm(receta_nombre))
                if receta is None:
                    raise ValueError(f"la receta '{receta_nombre}' no existe (revisa la hoja Recetas)")
                insumo = insumos_por_nombre.get(_norm(insumo_nombre))
                if insumo is None:
                    raise ValueError(f"el insumo '{insumo_nombre}' no existe (revisa la hoja Insumos)")
                if cantidad is None:
                    raise ValueError("falta la cantidad usada")

                db.add(RecetaInsumo(receta_id=receta.id, insumo_id=insumo.id, cantidad_usada=float(cantidad)))
                resultado["receta_insumos_creados"] += 1
            except Exception as exc:
                error("Receta_Insumos", fila_num, str(exc))
        db.flush()

    if "Ventas" in wb.sheetnames:
        sheet = wb["Ventas"]
        headers = _header_map(sheet)

        costo_actual_por_receta: dict[int, float] = {}
        for receta_insumo in db.query(RecetaInsumo).all():
            costo_actual_por_receta[receta_insumo.receta_id] = costo_actual_por_receta.get(
                receta_insumo.receta_id, 0.0
            ) + float(receta_insumo.cantidad_usada) * float(receta_insumo.insumo.costo_por_unidad_base)

        for fila_num, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            if _fila_vacia(row):
                continue
            try:
                receta_nombre = _get(row, headers, "receta")
                cantidad = _get(row, headers, "cantidad vendida", "cantidad_vendida")
                fecha_raw = _get(row, headers, "fecha")
                medio_pago = _norm(_get(row, headers, "medio de pago", "medio_pago"))
                total_venta_raw = _get(row, headers, "total de venta", "total_venta")
                costo_snapshot_raw = _get(row, headers, "costo de insumos", "costo_insumos_snapshot")

                if not receta_nombre:
                    raise ValueError("falta la receta")
                receta = recetas_por_nombre.get(_norm(receta_nombre))
                if receta is None:
                    raise ValueError(f"la receta '{receta_nombre}' no existe (revisa la hoja Recetas)")
                if cantidad is None:
                    raise ValueError("falta la cantidad vendida")
                if medio_pago not in MEDIOS_PAGO_VALIDOS:
                    raise ValueError(f"medio de pago invalido: '{medio_pago}' (usa efectivo o banco)")

                fecha = _parse_fecha(fecha_raw)
                cantidad_num = float(cantidad)
                precio_venta = float(receta.precio_venta)

                total_venta = (
                    float(total_venta_raw) if total_venta_raw not in (None, "") else cantidad_num * precio_venta
                )
                costo_snapshot = (
                    float(costo_snapshot_raw)
                    if costo_snapshot_raw not in (None, "")
                    else cantidad_num * costo_actual_por_receta.get(receta.id, 0.0)
                )

                db.add(
                    Venta(
                        receta_id=receta.id,
                        cantidad_vendida=int(cantidad_num),
                        fecha=fecha,
                        medio_pago=medio_pago,
                        total_venta=total_venta,
                        costo_insumos_snapshot=costo_snapshot,
                    )
                )
                resultado["ventas_creadas"] += 1
            except Exception as exc:
                error("Ventas", fila_num, str(exc))

    if "Gastos_Fijos" in wb.sheetnames:
        sheet = wb["Gastos_Fijos"]
        headers = _header_map(sheet)
        for fila_num, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            if _fila_vacia(row):
                continue
            try:
                concepto = _get(row, headers, "concepto")
                monto = _get(row, headers, "monto mensual", "monto_mensual")
                categoria = _get(row, headers, "categoria")
                medio_pago = _norm(_get(row, headers, "medio de pago", "medio_pago"))
                iva_raw = _get(row, headers, "iva", "tratamiento_fiscal", "tratamiento fiscal")
                mes_raw = _get(row, headers, "mes", "periodo")

                if not concepto:
                    raise ValueError("falta el concepto")
                if monto is None:
                    raise ValueError("falta el monto mensual")
                if not categoria:
                    raise ValueError("falta la categoria")
                if medio_pago not in MEDIOS_PAGO_VALIDOS:
                    raise ValueError(f"medio de pago invalido: '{medio_pago}' (usa efectivo o banco)")
                tratamiento_fiscal = _parse_tasa_iva(iva_raw)
                # si no se especifica el mes, se asume que es el gasto del mes en curso
                periodo = _parse_periodo(mes_raw) if mes_raw not in (None, "") else _parse_periodo(date.today())

                db.add(
                    GastoFijo(
                        concepto=str(concepto).strip(),
                        periodo=periodo,
                        monto_mensual=float(monto),
                        categoria=str(categoria).strip(),
                        medio_pago=medio_pago,
                        tratamiento_fiscal=tratamiento_fiscal,
                    )
                )
                resultado["gastos_fijos_creados"] += 1
            except Exception as exc:
                error("Gastos_Fijos", fila_num, str(exc))

    db.commit()
    return resultado
